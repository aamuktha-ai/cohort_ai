import sys
from pathlib import Path

from openpyxl import load_workbook

if len(sys.argv) != 2:
    raise SystemExit("Usage: python tools/inspect_workbook.py path/to/workbook.xlsx")

path = Path(sys.argv[1])
workbook = load_workbook(path, read_only=True, data_only=True)

print(workbook.sheetnames)

for worksheet in workbook.worksheets:
    print(f"### {worksheet.title} rows={worksheet.max_row} cols={worksheet.max_column}")
    for row in worksheet.iter_rows(min_row=1, max_row=min(worksheet.max_row, 8), values_only=True):
        values = ["" if value is None else str(value).replace("\n", " ")[:80] for value in row[:12]]
        print(" | ".join(values))
    print()
